# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n===== SHEET:", ws.title, " dims=", ws.dimensions, " max_row=", ws.max_row, " max_col=", ws.max_column, "=====")
    for row in ws.iter_rows():
        vals = []
        for c in row:
            if c.value is not None:
                vals.append(f"{c.coordinate}={c.value!r}")
        if vals:
            print(" | ".join(vals))
