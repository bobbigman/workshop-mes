# -*- coding: utf-8 -*-
"""生成《打蛋器生产工艺作业指导书》docx — 专业文书格式."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析\打蛋器_作业指导书.docx"

HEI = "黑体"; SONG = "宋体"; TIMES = "Times New Roman"

def set_font(run, cn=SONG, en=TIMES, size=12, bold=False, color=RGBColor(0,0,0)):
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def para(doc, text, cn=SONG, size=12, bold=False, align=None, indent=True,
         before=0, after=0, line=1.5, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    if line == 1.5: pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    else: pf.line_spacing = line
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent: pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    set_font(r, cn=cn, size=size, bold=bold)
    return p

def heading(doc, text, level):
    styles = {1: (16, 14, 6, HEI), 2: (14, 11, 5, HEI), 3: (12, 9, 4, HEI)}
    size, before, after, cn = styles.get(level, styles[2])
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    set_font(r, cn=cn, size=size, bold=True)
    h.paragraph_format.space_before = Pt(before)
    h.paragraph_format.space_after = Pt(after)
    return h

def title(doc, text, sub=None):
    t = doc.add_paragraph(style='Title')
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(text); set_font(r, cn=HEI, size=18, bold=True)
    t.paragraph_format.space_after = Pt(18)
    if sub:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = s.add_run(sub); set_font(rs, cn="楷体", size=12)
        s.paragraph_format.space_after = Pt(12)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_no_indent(cell):
    for p in cell.paragraphs:
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(p.text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT

def make_table(doc, rows, cols, header, data, widths=None):
    t = doc.add_table(rows=1+len(data), cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(header):
        c = t.cell(0, i); c.text = htxt
        set_cell_bg(c, "D9D9D9"); cell_no_indent(c)
        for p in c.paragraphs:
            for r in p.runs: set_font(r, cn=SONG, size=10.5, bold=True)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ri, row in enumerate(data, start=1):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = str(val)
            cell_no_indent(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                for r in p.runs: set_font(r, cn=SONG, size=10.5)
    # cantSplit on all rows + repeat header
    for tr in t.rows:
        trPr = tr._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
    tblHeader = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(tblHeader)
    return t

doc = Document()
# page: A4, margins 2.5cm
sec = doc.sections[0]
sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
for m in ('top_margin','bottom_margin','left_margin','right_margin'):
    setattr(sec, m, Cm(2.5))
# default style font
st = doc.styles['Normal']; st.font.name = TIMES; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

# ===== 标题页 =====
title(doc, "打蛋器生产工艺作业指导书", "厨具五金厂 CJ001 · 打蛋器 DDQ-001")
para(doc, "材质：304 不锈钢　　单位：件　　版本：V1.0　　生效日期：2026-09-25", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
para(doc, "本指导书依据厨具五金厂（CJ001）现行工艺路线编制，用于指导打蛋器各工序的作业、自检与质量控制。工序工艺参数为参考建议值，具体以现场实际工艺文件为准。")

# ===== 1 适用范围 =====
heading(doc, "1  适用范围与产品信息", 1)
para(doc, "本指导书适用于厨具五金厂打蛋器（产品编号 DDQ-001）从下料、冲压、打磨、质检、组装到包装的全部生产作业环节，供一线操作工、班组长和质检员使用。")
make_table(doc, 2, 2, ["项目", "内容"], [["产品编号/名称", "DDQ-001 / 打蛋器"], ["材质", "304 不锈钢"]], widths=None)

# ===== 2 工艺路线 =====
heading(doc, "2  工艺路线总览", 1)
para(doc, "打蛋器按“下料 → 冲压 → 打磨 → 质检 → 组装 → 包装”六道工序顺序流转，前道工序合格后方可流入下道工序。")
make_table(doc, 2, 3, ["序号", "工序", "主要不良品项"],
           [["1~6", "下料 → 冲压 → 打磨 → 质检 → 组装 → 包装", "毛刺 / 划伤 / 组装松动 / 包装破损"]])

# ===== 3 工序作业指导 =====
heading(doc, "3  各工序作业指导", 1)
OPS = [
    ("3.1  下料（XC）", "按图纸尺寸在 304 不锈钢板材上划线、裁切打蛋器所需坯料。", None,
     [["板材厚度", "1.0 mm（参考）"], ["裁切尺寸", "按图纸，留加工余量"]],
     "坯料尺寸符合图纸公差，无缺角、无料厚异常。", "下料无专门不良品项，尺寸超差按返工处理。"),
    ("3.2  冲压（CY）", "将下料坯料放入冲压模具，一次/多次冲压成型打蛋器手柄与打蛋头。", "毛刺",
     [["模具", "打蛋器成型模"], ["冲压次数", "按工艺要求（参考）"], ["毛刺控制", "成型后去毛刺"]],
     "成型尺寸合格，轮廓清晰，无明显变形。", "毛刺需去除到位，残留毛刺为不良。"),
    ("3.3  打磨（DM）", "对冲压件表面进行去毛刺、抛光处理，消除刀痕与毛刺。", "划伤",
     [["打磨方式", "砂轮/抛光"], ["表面要求", "光滑无明显划痕"]],
     "表面平整光滑，无明显划伤、氧化斑。", "打磨不当产生划伤为不良。"),
    ("3.4  质检（ZJ）", "对半成品进行尺寸、外观、装配性检验。", None,
     [["检验项", "尺寸 / 外观 / 装配"], ["抽检", "按批次抽检"]],
     "尺寸符合图纸，外观无缺陷，装配顺畅。", "质检无专门不良品项，不合格按返工或报废。"),
    ("3.5  组装（ZZ）", "将手柄与打蛋头等零件按图纸装配成成品。", "组装松动",
     [["装配", "手柄 + 打蛋头"], ["紧固", "牢固无松动"]],
     "装配到位，连接牢固，无松脱。", "组装松动为不良，须返工。"),
    ("3.6  包装（BZ）", "成品清洁后装入内衬与外箱，粘贴产品标识。", "包装破损",
     [["包装", "内衬 + 外箱"], ["标识", "产品信息 / 数量"]],
     "包装完好，标识清晰，数量准确。", "包装破损为不良。"),
]
for (h, content, defect, params, std, notice) in OPS:
    heading(doc, h, 2)
    para(doc, "作业内容：" + content)
    if defect:
        para(doc, "本工序关联不良品项：%s。" % defect)
    make_table(doc, len(params)+1, 2, ["工艺参数", "要求（参考）"], params)
    para(doc, "质量标准：" + std)
    para(doc, "注意：" + notice)

# ===== 4 通用要求 =====
heading(doc, "4  通用质量与安全要求", 1)
for line in [
    "操作前穿戴劳保用品，按设备操作规程作业，严禁违章操作。",
    "各工序自检合格后再报工；发现不良及时上报班组长。",
    "不良品与合格品分区放置，标识清楚，防止混料。",
    "生产数据以报工系统记录为准，报工数量须与实物一致。",
]:
    para(doc, line)

# ===== 5 附录 =====
heading(doc, "5  附录：不良品项对照表", 1)
make_table(doc, 2, 3, ["不良品项", "所属工序", "处理方式"],
           [["毛刺", "冲压", "返工去除"], ["划伤", "打磨", "返工抛光"], ["组装松动", "组装", "返工紧固"], ["包装破损", "包装", "返工更换包装"]])

doc.save(OUT)
print("saved:", OUT, "exists:", os.path.exists(OUT))
