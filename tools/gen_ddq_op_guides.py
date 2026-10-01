# -*- coding: utf-8 -*-
"""批量生成 6 份工序作业指导书 docx（下料/冲压/打磨/质检/组装/包装）."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析"
HEI = "黑体"; SONG = "宋体"; TIMES = "Times New Roman"

def set_font(run, cn=SONG, en=TIMES, size=12, bold=False, color=RGBColor(0, 0, 0)):
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color

def para(doc, text, size=12, bold=False, align=None, indent=True, before=0, after=0):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent: pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text); set_font(r, size=size, bold=bold)
    return p

def heading(doc, text, level):
    styles = {1: (16, 14, 6), 2: (14, 11, 5)}
    size, before, after = styles.get(level, styles[2])
    h = doc.add_heading("", level=level)
    r = h.add_run(text); set_font(r, cn=HEI, size=size, bold=True)
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
    return h

def title(doc, text, sub=None):
    t = doc.add_paragraph(style='Title'); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(text); set_font(r, cn=HEI, size=18, bold=True)
    t.paragraph_format.space_after = Pt(18)
    if sub:
        s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = s.add_run(sub); set_font(rs, cn="楷体", size=12)
        s.paragraph_format.space_after = Pt(12)

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_fmt(cell):
    for p in cell.paragraphs:
        p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.left_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(p.text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT

def make_table(doc, header, data):
    t = doc.add_table(rows=1 + len(data), cols=len(header))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(header):
        c = t.cell(0, i); c.text = htxt
        set_cell_bg(c, "D9D9D9"); cell_fmt(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs: set_font(r, size=10.5, bold=True)
    for ri, row in enumerate(data, start=1):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = str(val)
            cell_fmt(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                for r in p.runs: set_font(r, size=10.5)
    for tr in t.rows:
        tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t

def build(op_name, op_code, content, steps, params, quality, defect, safety, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    for m in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, m, Cm(2.5))
    st = doc.styles['Normal']; st.font.name = TIMES; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

    title(doc, "%s工序作业指导书" % op_name, "厨具五金厂 CJ001 · 打蛋器 DDQ-001 工艺路线（%s %s）" % (op_code, op_name))
    para(doc, "本指导书用于指导打蛋器%s工序的作业、自检与质量控制。工序工艺参数为参考建议值，具体以现场实际工艺文件为准。" % op_name, after=6)

    heading(doc, "1  适用范围与作业内容", 1)
    para(doc, content)

    heading(doc, "2  操作步骤", 1)
    for i, s in enumerate(steps, 1):
        para(doc, "%d. %s" % (i, s), indent=False)

    heading(doc, "3  工艺参数（参考）", 1)
    make_table(doc, ["参数", "要求（参考）"], params)

    heading(doc, "4  质量标准", 1)
    para(doc, quality)

    heading(doc, "5  不良品项", 1)
    para(doc, defect)

    heading(doc, "6  安全注意事项", 1)
    para(doc, safety)

    doc.save(out)
    print("saved:", os.path.basename(out))

OPS = [
    ("下料", "XC", "按图纸尺寸在 304 不锈钢板材上划线、裁切，获得打蛋器所需坯料。",
     ["领取板材，核对材质、规格与批次", "按图纸划线，标注下料尺寸", "裁切坯料，保证尺寸符合公差", "自检坯料外观，报工并转下道工序"],
     [["板材厚度", "1.0 mm（参考）"], ["裁切尺寸", "按图纸，留加工余量"], ["公差", "按图纸要求"]],
     "坯料尺寸符合图纸公差，无缺角、无料厚异常、无过烧变色。",
     "下料无专门不良品项；尺寸超差不合格按返工处理。",
     "操作设备须持证上岗，佩戴防护用品，严禁戴手套操作旋转部位。"),
    ("冲压", "CY", "将下料坯料放入冲压模具，冲压成型打蛋器手柄与打蛋头。",
     ["安装并校核模具", "首件试冲，确认成型合格", "批量冲压，中途抽检", "去毛刺，自检后报工"],
     [["模具", "打蛋器成型模"], ["冲压次数", "按工艺要求（参考）"], ["毛刺控制", "成型后去毛刺"]],
     "成型尺寸合格，轮廓清晰，无明显变形，无开裂。",
     "关联不良品项：毛刺。残留毛刺须返工去除。",
     "严禁手入模区；设备运行时不得调整模具；两人以上操作须统一口令。"),
    ("打磨", "DM", "对冲压件进行去毛刺与抛光处理，消除刀痕、毛刺。",
     ["领取冲压件", "砂轮/抛光去毛刺", "抛光表面至要求", "检查表面质量，报工"],
     [["打磨方式", "砂轮 / 抛光"], ["表面要求", "光滑，无明显划痕"]],
     "表面平整光滑，无毛刺、无明显划伤、无氧化斑。",
     "关联不良品项：划伤。打磨不当造成划伤为不良，须返工抛光。",
     "佩戴护目镜与防护手套；抛光粉尘区须通风；设备停机后清理。"),
    ("质检", "ZJ", "对半成品进行尺寸、外观与装配性检验，判定是否合格。",
     ["按批次抽取样件", "测量关键尺寸", "检查外观缺陷", "试装配检验，记录并报工"],
     [["检验项", "尺寸 / 外观 / 装配"], ["抽检比例", "按批次抽检（参考）"], ["量具", "卡尺、检具"]],
     "尺寸符合图纸，外观无缺陷，装配顺畅，检验记录完整。",
     "质检无专门不良品项；不合格件按返工或报废处理，并登记不良原因。",
     "量具轻拿轻放、定期校准；不合格品分区放置并标识。"),
    ("组装", "ZZ", "将手柄与打蛋头等零件按图纸装配成成品。",
     ["领取已检零件", "按图纸装配手柄与打蛋头", "紧固并检验牢固度", "自检后报工"],
     [["装配关系", "手柄 + 打蛋头"], ["紧固要求", "牢固无松动"]],
     "装配到位，连接牢固，无松脱，外观无损伤。",
     "关联不良品项：组装松动。松动须返工紧固。",
     "装配工位保持整洁；使用装配工具时防止夹伤；不良件及时隔离。"),
    ("包装", "BZ", "成品清洁后装入内衬与外箱，粘贴产品标识并入库。",
     ["清洁成品表面", "装入内衬与外箱", "粘贴产品标识与数量", "封箱、码放、入库"],
     [["包装", "内衬 + 外箱"], ["标识", "产品信息 / 数量"], ["码放", "按库位要求"]],
     "包装完好，标识清晰，数量准确，码放整齐。",
     "关联不良品项：包装破损。破损须返工更换包装。",
     "重箱轻放，码放不超高；封箱胶带使用注意安全；库位标识对齐。"),
]

for (op_name, op_code, content, steps, params, quality, defect, safety) in OPS:
    out = os.path.join(ROOT, "%s工序作业指导书.docx" % op_name)
    build(op_name, op_code, content, steps, params, quality, defect, safety, out)
print("ALL DONE")
