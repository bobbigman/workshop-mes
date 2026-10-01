# -*- coding: utf-8 -*-
"""阶段 D：回填 docs/AI知识库检索命中基准表.xlsx 验收状态与备注（原路径原地更新）。"""
import sys
from pathlib import Path

sys.path.insert(0, r'C:\Users\bob\AppData\Local\Qianwen\User Data\qwen-agent\HOeGDzsMN3\skills\excel\scripts')
from formula_verify import task_checklist, save_verified
from openpyxl import load_workbook

work_dir = Path(r'D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析\kb-task\excel-20260924T155212+0800-a8_qrbha')
task_token = work_dir.name
source = Path(r'D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析\docs\AI知识库检索命中基准表.xlsx')
out = source

# (行, 序号, 知识域, 问法, 预期命中, 覆盖状态, 依据/来源) —— 与源表逐字一致，用于防误改核验
ROWS = [
    (3, 1, '01-操作入门', '怎么登录系统？都有哪些角色？', '01-操作入门.md #登录/角色', '现有', 'SupportDocs/README 测试问法 #4'),
    (4, 2, '01-操作入门', '第一次使用先配置什么？顺序是什么？', '01-操作入门.md #第一次配置', '现有', '01-操作入门.md'),
    (5, 3, '01-操作入门', '怎么给工人建账号？怎么改密码？', '01-操作入门.md #建账号/改密码', '现有', 'SupportDocs/README 测试问法 #4'),
    (6, 4, '01-操作入门', '怎么建部门、把人加进部门？', '01-操作入门.md #建部门', '现有', '01-操作入门.md'),
    (7, 5, '02-基础资料', '怎么配单位/不良品项/工序？', '02-基础资料.md #单位/不良品项/工序', '现有', '02-基础资料.md'),
    (8, 6, '02-基础资料', '工序报工权限怎么配？', '02-基础资料.md #工序（报工权限）', '现有', 'SupportDocs/README 测试问法 #5'),
    (9, 7, '02-基础资料', '怎么配工艺路线和产品？', '02-基础资料.md #工艺路线/产品', '现有', '02-基础资料.md'),
    (10, 8, '02-基础资料', '产品编号可以重复吗？', '02-基础资料.md #产品编号唯一', '现有', '02-基础资料.md'),
    (11, 9, '03-工单', '怎么下工单？', '03-工单.md #创建工单', '现有', 'SupportDocs/README 测试问法 #3'),
    (12, 10, '03-工单', '怎么打印二维码流转卡？', '03-工单.md #打印', '现有', 'SupportDocs/README 测试问法 #12'),
    (13, 11, '03-工单', '工单有哪些状态？怎么结束/撤回/取消？', '03-工单.md #状态流转', '现有', 'SupportDocs/README 测试问法 #11'),
    (14, 12, '03-工单', '怎么查进度、报表、看板？', '03-工单.md #查进度/报表/看板', '现有', '03-工单.md'),
    (15, 13, '04-报工与复核', '怎么电脑报工？', '04-报工与复核.md #电脑报工', '现有', '04-报工与复核.md'),
    (16, 14, '04-报工与复核', '怎么补报？一键补报未完工序怎么用？', '04-报工与复核.md #补报', '现有', 'SupportDocs/README 测试问法 #1'),
    (17, 15, '04-报工与复核', '报工复核在哪操作？怎么退回？', '04-报工与复核.md #复核', '现有', 'SupportDocs/README 测试问法 #7'),
    (18, 16, '04-报工与复核', '报错了怎么改数？', '04-报工与复核.md #改报工数', '现有', 'SupportDocs/README 测试问法 #8'),
    (19, 17, '04-报工与复核', '手机怎么扫码报工？PDA 怎么用？', '04-报工与复核.md #扫码/#PDA', '现有', '04-报工与复核.md'),
    (20, 18, '05-常见问题', '为什么产品删不掉？', '05-常见问题.md #删除顺序', '现有', 'SupportDocs/README 测试问法 #2'),
    (21, 19, '05-常见问题', '无报工权限是什么意思？', '05-常见问题.md #无报工权限', '现有', 'SupportDocs/README 测试问法 #6'),
    (22, 20, '05-常见问题', '网页打不开/登录失败怎么办？', '05-常见问题.md #网页打不开', '现有', '05-常见问题.md'),
    (23, 21, '05-常见问题', '补报和改数有什么区别？', '05-常见问题.md #补报vs改数', '现有', '05-常见问题.md'),
    (24, 22, '05-常见问题', 'AI 不能用/页面报服务器错误？', '05-常见问题.md #AI不能用', '现有', '05-常见问题.md'),
    (25, 23, '06-数据口径', '报工后为什么生产报表没增加？', '06-数据口径.md #双轨口径', '现有', 'SupportDocs/README 测试问法 #10'),
    (26, 24, '06-数据口径', '完成数为什么是 80 不是 180？', '06-数据口径.md #完成数', '现有', 'SupportDocs/README 测试问法 #9'),
    (27, 25, '06-数据口径', '工单找不到/结束后又变执行中？', '06-数据口径.md #工单状态', '现有', '06-数据口径.md'),
    (28, 26, '06-数据口径', '超出计划数/剩余可报为零？', '06-数据口径.md #超计划数', '现有', '06-数据口径.md'),
    (29, 27, '工资与工价', '工资怎么算？工价表在哪配？', '工资与工价（docs/30 生成篇）', '新增', 'docs/21；61 §2.1'),
    (30, 28, '工资与工价', '生成的工资单为什么确认不了？', '工资与工价 #未匹配标红阻断', '新增', 'docs/21；61 §2.1'),
    (31, 29, '工资与工价', '工资单怎么生成/确认/导出？', '工资与工价 #工资报表', '新增', 'docs/21；61 §2.1'),
    (32, 30, '派工到人', '怎么把工序派给工人？', '派工到人 #工单详情执行人', '新增', 'docs/29；61 §2.2'),
    (33, 31, '派工到人', '我的任务在哪看？', '派工到人 #工人「我的任务」', '新增', 'docs/29；61 §2.2'),
    (34, 32, '异常上报', '工人怎么上报异常？', '异常上报 #H5上报', '新增', 'docs/28；61 §2.3'),
    (35, 33, '异常上报', '异常处理在哪弄？看板红标怎么消？', '异常上报 #PC处理回填', '新增', 'docs/28；61 §2.3'),
    (36, 34, '执行监控', '执行监控看什么？', '执行监控（docs/30 生成篇）', '新增', 'docs/19 补充执行；61 §2.4'),
    (37, 35, '执行监控', '工序进度怎么算？为什么跟报表对不上？', '执行监控 #进度口径', '新增', 'docs/19 补充执行；61 §2.4'),
    (38, 36, '批量导入', '批量导入怎么用？', '批量导入 #下载母版/上传', '新增', 'docs/57；61 §2.5'),
    (39, 37, '批量导入', '导入提示第几行失败怎么办？', '批量导入 #行级错误', '新增', 'docs/57；61 §2.5'),
    (40, 38, '多账套选厂', '登录为什么要选厂？', '多账套选厂 #登录选厂', '新增', 'docs/52；61 §2.6'),
    (41, 39, '多账套选厂', '怎么切到另一家厂？', '多账套选厂 #工厂隔离/切换', '新增', 'docs/52；61 §2.6'),
    (42, 40, '微信预警', '微信预警怎么开？', '微信预警 #默认关闭需配置', '新增', 'docs/20；61 §2.7'),
]

BASE_H = {8: '三条硬规则之一', 10: '三条硬规则之一', 20: '三条硬规则之一', 37: '口径题重点', 42: '低优先级'}
GEN = '；2026-09-24 实测已命中（按检索服务口径复测，预期章节进入 Top 结果）'
SUFFIX = {r: GEN for r in range(29, 43)}
SUFFIX[18] = '；Top2 命中预期 04#报错数量怎么改，Top1 为 05#报错了不确定是否提交成功（同主题，两条同入检索证据，回答正确）'
for r in (25, 26, 27, 28):
    SUFFIX[r] = '；05 与 06 同名章节并列 Top 结果，预期 06 章节命中'
EXPECTED_H = {r: ((BASE_H.get(r) or '') + (SUFFIX.get(r) or '')) or None for r in range(3, 43)}


def build_workbook():
    wb = load_workbook(source, data_only=False)
    ws = wb['命中基准']
    for r in range(3, 43):
        ws.cell(r, 6).value = '已命中'
    for r, remark in EXPECTED_H.items():
        if remark:
            ws.cell(r, 8).value = remark
    return wb


def r1(v, f):
    ws = v['命中基准']
    return all(ws.cell(r, 6).value == '已命中' for r in range(3, 43))


def r2(v, f):
    ws = v['命中基准']
    return all((ws.cell(r, 8).value or None) == EXPECTED_H[r] for r in range(3, 43))


def r3(v, f):
    ws = v['命中基准']
    ok = (ws['A1'].value == 'AI 知识库（SupportDocs）检索命中基准表 —— 现有六篇 + 61 补全七项'
          and [ws.cell(2, c).value for c in range(1, 9)] ==
          ['序号', '知识域', '问法（用户怎么说）', '预期命中（文档/章节）', '覆盖状态', '验收状态', '依据/来源', '备注'])
    for r, no, dom, q, e, st, src in ROWS:
        ok = ok and (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value,
                     ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 7).value) == (no, dom, q, e, st, src)
    return ok


def r4(v, f):
    wsf = f['验收汇总']
    labels = ['总问法数', '现有六篇覆盖条数', '新增功能补全条数', '已命中', '未命中', '待验收', '通过率（已命中/总）']
    ok = all(wsf.cell(r, 1).value == t for r, t in zip(range(3, 10), labels))
    ok = ok and all(isinstance(wsf.cell(r, 2).value, str) and wsf.cell(r, 2).value.startswith('=') for r in range(3, 10))
    wsv = v['验收汇总']
    ok = ok and wsv.cell(13, 1).value == '1. 逐条用问法向检索接口 / 扣子 MCP 实测，命中预期章节则把验收状态改为『已命中』，未命中改『未命中』，未测保持『待验收』。'
    b6 = wsv.cell(6, 2).value
    if isinstance(b6, (int, float)):
        ok = ok and b6 == 40
    return ok


requirements = [
    {'id': 'R1', 'quote': '40 条问法验收状态回填', 'check': '命中基准 F3:F42 全部为『已命中』（本次实测全部命中，无未命中项）'},
    {'id': 'R2', 'quote': '备注写明实测说明，不得静默', 'check': '备注列 H 等于原定内容：原有条目备注（三条硬规则之一/口径题重点/低优先级）保留并按时追加实测说明；#16、#23～26 附特殊说明；其余现有条备注保持空'},
    {'id': 'R3', 'quote': '不破坏问法数据', 'check': '标题行、表头行与 A/C/D/E/G 列 40 行数据逐格与源一致（仅 F、H 列被更新）'},
    {'id': 'R4', 'quote': '验收汇总公式不手填', 'check': '验收汇总 B3:B9 仍为公式；标签与使用说明文字未被改动；若有重算缓存则已命中=40'},
]

checks = {'R1': r1, 'R2': r2, 'R3': r3, 'R4': r4}

if __name__ == '__main__':
    verify = task_checklist(work_dir, task_token, checks, requirements=requirements, outputs=(out,))
    wb = build_workbook()
    receipt = save_verified(wb, out, verify)
    print({k: receipt.get(k) for k in ('output', 'checks_passed', 'total_formulas', 'total_errors')})
