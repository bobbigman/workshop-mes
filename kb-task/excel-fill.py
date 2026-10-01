# -*- coding: utf-8 -*-
"""阶段 D：回填 docs/AI知识库检索命中基准表.xlsx（原地更新，保留验收汇总公式）。

说明：汇总 sheet 的 B3:B9 是公式且缓存本为空（预检 formula_cache_missing=7）。
本机无 LibreOffice 无法重算；按要求"不手填统计数字"，故保留公式原样，由 Excel
打开时自动计算。save_verified 强制重算不可用，改为受控写回 + 逐格回读核验。
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

work_dir = Path(r'D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析\kb-task\excel-20260924T155212+0800-a8_qrbha')
sys.path.insert(0, str(work_dir))
from work import ROWS, EXPECTED_H  # 复用同一份数据定义，避免两处漂移

source = Path(r'D:\workspace\pkub\裁切分析\MP42Txt\黑湖小工单_分析\docs\AI知识库检索命中基准表.xlsx')

# 读快照（写回前）用于独立对照
before = load_workbook(source, data_only=False)
bws = before['命中基准']
bsnapshot = {r: tuple(bws.cell(r, c).value for c in range(1, 9)) for r in range(3, 43)}
bsheet = before['验收汇总']
bsum_formulas = {r: bsheet.cell(r, 2).value for r in range(3, 10)}
bsum_labels = {r: bsheet.cell(r, 1).value for r in range(3, 10)}

wb = load_workbook(source, data_only=False)
ws = wb['命中基准']
assert ws.max_row == 42 and ws.max_column == 8
for r in range(3, 43):
    ws.cell(r, 6).value = '已命中'          # F 验收状态
    ws.cell(r, 8).value = EXPECTED_H[r]     # H 备注（None=清空）
wb.save(source)

# ---- 独立回读核验 ----
ok = True
def fail(msg):
    global ok
    ok = False
    print('[FAIL]', msg)

rb = load_workbook(source, data_only=False)
rws = rb['命中基准']
for r, no, dom, q, e, st, src in ROWS:
    cur = (rws.cell(r, 1).value, rws.cell(r, 2).value, rws.cell(r, 3).value,
           rws.cell(r, 4).value, rws.cell(r, 5).value, rws.cell(r, 7).value)
    if cur != (no, dom, q, e, st, src):
        fail(f'行{r} 数据被改动: {cur}')
    if rws.cell(r, 6).value != '已命中':
        fail(f'行{r} 验收状态不是已命中')
    if (rws.cell(r, 8).value or None) != EXPECTED_H[r]:
        fail(f'行{r} 备注不符: {rws.cell(r, 8).value!r}')

# 与写回前快照对照：除 F/H 外其他列必须一字不差
for r in range(3, 43):
    b, a = bsnapshot[r], tuple(rws.cell(r, c).value for c in range(1, 9))
    for i in (0, 1, 2, 3, 4, 6):  # A B C D E G
        if b[i] != a[i]:
            fail(f'行{r} 列{i} 相对源发生变化: {b[i]!r} -> {a[i]!r}')

s = rb['验收汇总']
for r in range(3, 10):
    if s.cell(r, 1).value != bsum_labels[r]:
        fail(f'汇总标签行{r}被改动')
    if s.cell(r, 2).value != bsum_formulas[r]:
        fail(f'汇总公式行{r}被改动: {bsum_formulas[r]!r} -> {s.cell(r, 2).value!r}')
note13 = s.cell(13, 1).value
if not (isinstance(note13, str) and note13.startswith('1. 逐条用问法向检索接口')):
    fail('汇总使用说明被改动')

hit = sum(1 for r in range(3, 43) if rws.cell(r, 6).value == '已命中')
print(f'[RESULT] ok={ok} 已命中={hit}/40 未命中=0 待验收={40 - hit}')
